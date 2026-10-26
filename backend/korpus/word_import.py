"""Bounded XLSX parsing; import creates drafts and never changes existing entries."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError

from django.core.exceptions import ValidationError
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .admin_forms import EquivalentsWidget
from .models import WordEntry
from .reading import normalize_word

LANGUAGES = [code for code, _ in EquivalentsWidget.languages]
HEADERS = ['word', 'meaning', 'source', 'pronunciation', 'example', 'status', *LANGUAGES]
MAX_ROWS = 1000


def template_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = 'So‘zlar'
    sheet.append(HEADERS)
    sheet.freeze_panes = 'A2'
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    instructions = book.create_sheet('Yo‘riqnoma')
    for line in [
        'Birinchi varaqda har qatorga bitta so‘z yozing. Sarlavhalarni o‘zgartirmang.',
        'word — so‘z, meaning — ma’nosi, source — manba: majburiy.',
        'pronunciation — talaffuz, example — misol: ixtiyoriy.',
        'status: historical yoki community; bo‘sh bo‘lsa historical.',
        'uz, tr, kk va boshqa til ustunlariga shu tildagi muqobil so‘zni yozing.',
        '1000 qatorgacha, 5 MB gacha. Faqat .xlsx. Formulalardan foydalanmang.',
        'So‘z va manbasi mos kelgan mavjud qatorlar o‘tkazib yuboriladi.',
        'Barcha yangi so‘zlar qoralama saqlanadi. Audio alohida yuklanadi.',
    ]:
        instructions.append([line])
    instructions.column_dimensions['A'].width = 110
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def entry_key(word, source):
    return normalize_word(word), source.strip().casefold()


def parse_workbook(upload):
    if not upload.name.lower().endswith('.xlsx') or upload.size > 5 * 1024 * 1024:
        raise ValidationError('5 MB gacha bo‘lgan .xlsx fayl tanlang.')
    book = None
    try:
        with ZipFile(upload) as archive:
            if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                raise ValidationError('Excel faylining ochilgan hajmi 20 MB dan oshmasin.')
        upload.seek(0)
        book = load_workbook(upload, read_only=True, data_only=False, keep_links=False)
        sheet = book.worksheets[0]
        # Do not trust worksheet dimensions supplied by the file author.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=len(HEADERS) + 1)
        header_cells = next(iterator, ())
        headers = [str(cell.value).strip() if cell.value is not None else '' for cell in header_cells]
        while headers and not headers[-1]:
            headers.pop()
        if (not {'word', 'meaning', 'source'}.issubset(headers)
                or len(set(headers)) != len(headers) or any(h not in HEADERS for h in headers)):
            raise ValidationError('Ustunlar noto‘g‘ri. Namuna faylini yuklab olib to‘ldiring.')
        entries, errors, text_size = [], [], 0
        for number, cells in enumerate(iterator, 2):
            if not any(cell.value is not None for cell in cells):
                continue
            if number > MAX_ROWS + 1:
                raise ValidationError('Bir faylda ko‘pi bilan 1000 qator bo‘lsin.')
            if any(cell.data_type == 'f' for cell in cells):
                errors.append(f'{number}-qator: formula o‘rniga oddiy matn kiriting.')
                continue
            if any(cell.value is not None for cell in cells[len(headers):]):
                errors.append(f'{number}-qator: sarlavhasiz ustunda ma’lumot bor.')
                continue
            values = {key: str(cell.value).strip() if cell.value is not None else ''
                      for key, cell in zip(headers, cells)}
            text_size += sum(len(value) for value in values.values())
            if text_size > 200000:
                raise ValidationError('Matn hajmi katta. Faylni kichik qismlarga ajratib yuklang (200 000 belgi).')
            data = {key: values.get(key, '') for key in ['word', 'meaning', 'source', 'pronunciation', 'example']}
            data['status'] = values.get('status') or 'historical'
            data['equivalents'] = {key: values[key] for key in LANGUAGES if values.get(key)}
            try:
                WordEntry(**data).full_clean()
                entries.append(data)
            except ValidationError as exc:
                errors.append(f'{number}-qator: {"; ".join(exc.messages)}')
        if errors:
            raise ValidationError(errors[:30])
        if not entries:
            raise ValidationError('Faylda import qilinadigan so‘z topilmadi.')
        return entries
    except (BadZipFile, InvalidFileException, ValueError, KeyError, OSError, IndexError, ParseError) as exc:
        raise ValidationError('Excel faylni o‘qib bo‘lmadi. Namuna asosida yangi .xlsx saqlang.') from exc
    finally:
        if book:
            book.close()


def partition_entries(entries):
    seen = {entry_key(word, source) for word, source in WordEntry.objects.values_list('word', 'source')}
    new, skipped = [], 0
    for data in entries:
        key = entry_key(data['word'], data['source'])
        if key in seen:
            skipped += 1
        else:
            seen.add(key)
            new.append(data)
    return new, skipped
